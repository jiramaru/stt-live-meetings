"""Transcript export formats."""

from __future__ import annotations

import io
from datetime import datetime


def _clock(seconds: float) -> str:
    seconds = int(seconds)
    return f"{seconds // 3600:02d}:{seconds % 3600 // 60:02d}:{seconds % 60:02d}"


def _srt_time(seconds: float) -> str:
    ms = int(round(seconds * 1000))
    return f"{ms // 3_600_000:02d}:{ms % 3_600_000 // 60_000:02d}:{ms % 60_000 // 1000:02d},{ms % 1000:03d}"


def speaker_name(session: dict, speaker: str | None) -> str | None:
    if not speaker:
        return None
    return session.get("speakers", {}).get(speaker) or f"Intervenant {speaker.lstrip('S')}"


def _turns(session: dict):
    """Segments with the speaker name, given only when the speaker changes."""
    previous = object()
    for s in session["segments"]:
        speaker = s.get("speaker")
        name = speaker_name(session, speaker) if speaker != previous else None
        previous = speaker
        yield s, name


def _started(session: dict) -> str:
    return datetime.fromisoformat(session["started_at"]).astimezone().strftime("%d/%m/%Y %H:%M")


def _notes(session: dict) -> str:
    return (session.get("notes") or "").strip()


def to_txt(session: dict) -> str:
    lines = [session["title"], f"Date : {_started(session)}", ""]
    if _notes(session):
        lines += ["Notes :", _notes(session), "", "Transcription :", ""]
    for s, name in _turns(session):
        if name:
            lines += ["", f"{name} :"] if lines[-1] else [f"{name} :"]
        lines.append(f"[{_clock(s['start'])}] {s['text']}")
    return "\n".join(lines) + "\n"


def to_markdown(session: dict) -> str:
    lines = [f"# {session['title']}", "", f"*Date : {_started(session)}*", ""]
    if _notes(session):
        lines += ["## Notes", "", _notes(session), "", "## Transcription", ""]
    for s, name in _turns(session):
        if name:
            lines += [f"### {name}", ""] if not lines[-1] else ["", f"### {name}", ""]
        lines.append(f"**{_clock(s['start'])}** {s['text']}  ")
    return "\n".join(lines) + "\n"


def to_srt(session: dict) -> str:
    blocks = []
    for i, s in enumerate(session["segments"], start=1):
        name = speaker_name(session, s.get("speaker"))
        text = f"{name} : {s['text']}" if name else s["text"]
        blocks.append(f"{i}\n{_srt_time(s['start'])} --> {_srt_time(s['end'])}\n{text}\n")
    return "\n".join(blocks)


def to_docx(session: dict) -> bytes:
    from docx import Document
    from docx.shared import Pt, RGBColor

    doc = Document()
    doc.add_heading(session["title"], level=1)
    doc.add_paragraph(f"Date : {_started(session)}")
    if _notes(session):
        doc.add_heading("Notes", level=2)
        for line in _notes(session).splitlines():
            doc.add_paragraph(line)
        doc.add_heading("Transcription", level=2)
    for s, name in _turns(session):
        if name:
            heading = doc.add_paragraph()
            heading.paragraph_format.space_before = Pt(10)
            heading.add_run(name).bold = True
        p = doc.add_paragraph()
        stamp = p.add_run(f"[{_clock(s['start'])}]  ")
        stamp.font.size = Pt(9)
        stamp.font.color.rgb = RGBColor(0x66, 0x66, 0x66)
        p.add_run(s["text"])
    out = io.BytesIO()
    doc.save(out)
    return out.getvalue()


# format -> (renderer, media type, extension)
EXPORTERS = {
    "txt": (to_txt, "text/plain; charset=utf-8", "txt"),
    "md": (to_markdown, "text/markdown; charset=utf-8", "md"),
    "srt": (to_srt, "application/x-subrip; charset=utf-8", "srt"),
    "docx": (
        to_docx,
        "application/vnd.openxmlformats-officedocument.wordprocessingml.document",
        "docx",
    ),
}
